# 仕事チャンネル 型A-1・A-3・C のデータ（手順書 §3 型A・型C・§8、編集方針の出典の決まり）
#   python -X utf8 pipeline\fetch_a1c_data.py              … 取って data\raw\estat\ と data\job\a1c_r7.json に保存（同じ月に取った物があれば取らない）
#   python -X utf8 pipeline\fetch_a1c_data.py --force      … 同じ月でも取り直す
#   python -X utf8 pipeline\fetch_a1c_data.py --register   … 取った後、data\sources.json と data\checksums.json に自分のキーだけ足す
#     （ほかのエージェントが同じ2ファイルを直す時期は、--register を付けずに走らせ、表示された内容を手で足す）
# ・取る物（どれも e-Stat のファイル。値は表の値そのまま）
#   雇用動向調査 令和7年 年計 第12表 … 産業（大分類）別の、転職入職者の「前の勤め先を辞めた理由」（A-1 の辞める理由）
#   賃金構造基本統計調査 令和7年 一般労働者 職種 第5表 … 職種（小分類）×年齢階級（C-3 の新人とベテラン・「〇〇の40年」の数え）
#   賃金構造基本統計調査 令和7年 一般労働者 産業大分類 第1表 … 産業計と産業ごとの平均（型A・C の「全産業の平均」と業界の平均）
# ・年収・時給の計算は fetch_job_data.derive と同じ式（年収＝きまって支給×12＋賞与、時給＝年収÷（月の労働時間×12））
import argparse, datetime, json, os, re, sys
import openpyxl, xlrd

sys.path.insert(0, os.path.dirname(__file__))
import estat
from fetch_job_data import FIELDS, UNITS, DL, derive, download, num, sha256

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "job", "a1c_r7.json")
TODAY = datetime.date.today().isoformat()

# 雇用動向調査 第12表の列（就業形態計・年齢計の段）。人数は3〜24列、構成比は25〜46列
YAMETA = ["仕事の内容に興味を持てず", "能力・個性・資格を生かせず", "職場の人間関係", "会社の将来が不安", "収入が少ない",
          "労働条件が悪い", "結婚", "出産・育児", "介護・看護", "その他の個人的理由", "定年、契約期間の満了", "会社都合",
          "その他の理由（出向等を含む）"]
ERANDA = ["仕事の内容に興味があった", "能力・個性・資格を生かせる", "会社の将来に期待できる", "収入が多い", "労働条件がよい",
          "通勤が便利", "とにかく仕事に就きたかった", "その他の理由（出向等を含む）"]
AGE9 = ["20～24歳", "25～29歳", "30～34歳", "35～39歳", "40～44歳", "45～49歳", "50～54歳", "55～59歳", "60～64歳"]


def find_file(stats_code, word, dataset_has, table_prefix):
    """カタログ API（検索語つき）で、データセット名に dataset_has を全部含み、表名が table_prefix で始まる Excel を1つ探す"""
    start = 1
    while True:
        d = estat.call("getDataCatalog", statsCode=stats_code, searchWord=word, limit=100, startPosition=start)
        r = d["GET_DATA_CATALOG"]
        ds = r.get("DATA_CATALOG_LIST_INF", {}).get("DATA_CATALOG_INF", [])
        ds = [ds] if isinstance(ds, dict) else ds
        for x in ds:
            dname = x["DATASET"]["TITLE"].get("NAME", "")
            if not all(w in dname for w in dataset_has):
                continue
            res = x["RESOURCES"]["RESOURCE"]
            res = [res] if isinstance(res, dict) else res
            for rr in res:
                if rr["TITLE"].get("NAME", "").startswith(table_prefix) and rr.get("FORMAT", "").startswith("XLS"):
                    sid = rr["URL"].split("statInfId=")[1][:12]
                    return {"statInfId": sid, "dataset": dname, "table": rr["TITLE"]["NAME"],
                            "release_date": rr.get("RELEASE_DATE"), "landing_page": x["DATASET"].get("LANDING_PAGE"),
                            "url": DL.format(sid), "licence": rr.get("RESOURCE_LICENCE_ID")}
        nxt = r.get("DATA_CATALOG_LIST_INF", {}).get("RESULT_INF", {}).get("NEXT_KEY")
        if not nxt:
            sys.exit(f"【止まりました】e-Stat に表が見つかりません：{dataset_has} {table_prefix}")
        start = nxt


def clean(s):
    return re.sub(r"\s", "", str(s))


def zen2han(s):
    return s.translate(str.maketrans("０１２３４５６７８９", "0123456789"))


def age_label(s):
    """「　２０　～　２４歳」→「20～24歳」、「～１９歳」→「～19歳」、「７０歳～」→「70歳～」"""
    return zen2han(clean(s))


# ---- 1. 雇用動向調査 第12表 ----
def parse_t12(path):
    ws = xlrd.open_workbook(path, on_demand=True).sheet_by_index(0)
    title = clean(ws.cell_value(2, 3)) + clean(ws.cell_value(3, 3))
    if "前の勤め先を辞めた理由" not in title or "前職雇用者" not in title:
        sys.exit("【止まりました】雇用動向調査 第12表の題が想定と違います")
    # 最初の段が「就業形態 計・年齢 計」か
    if clean(ws.cell_value(5, 3)) != "計" or clean(ws.cell_value(5, 8)) != "計":
        sys.exit("【止まりました】雇用動向調査 第12表の最初の段が「就業形態計・年齢計」ではありません")
    # 列の見出しを確かめる（セルの中の改行・空白は除いて比べる）
    def head(c):
        return "".join(clean(ws.cell_value(r, c)) for r in range(6, 10))
    for i, name in enumerate(YAMETA):
        if clean(name) not in head(4 + i):
            sys.exit(f"【止まりました】第12表の {4 + i} 列目が「{name}」ではありません（{head(4 + i)}）")
    for i, name in enumerate(ERANDA):
        if clean(name) not in head(17 + i):
            sys.exit(f"【止まりました】第12表の {17 + i} 列目が「{name}」ではありません（{head(17 + i)}）")
    out, sec = {}, None
    for i in range(11, ws.nrows):
        a, b = clean(ws.cell_value(i, 1)), clean(ws.cell_value(i, 2))
        if a == "経験1年未満":  # 次の段（年齢階級別）に入ったら終わり
            break
        m = re.fullmatch(r"計（(男女計|男|女)）", b)
        if a == "ＧＴ" and m:
            sec, key = m.group(1), "産業計"
        elif len(a) == 1 and "Ｃ" <= a <= "Ｒ":
            key = a  # 産業の記号（産業名は表の中で切れているので、記号で引く）
        else:
            continue  # 企業規模の行
        row = ws.row_values(i)
        if num(row[25]) is None or abs(row[25] - 100) > 0.05:
            sys.exit(f"【止まりました】第12表の {i} 行目の構成比の計が100ではありません")
        out.setdefault(sec, {})[key] = {
            "表の名前": b, "転職入職者数_千人": num(row[3]),
            "辞めた理由_千人": {n: num(row[4 + k]) for k, n in enumerate(YAMETA)},
            "辞めた理由_構成比": {n: num(row[26 + k]) for k, n in enumerate(YAMETA)},
            "選んだ理由_千人": {n: num(row[17 + k]) for k, n in enumerate(ERANDA)},
            "選んだ理由_構成比": {n: num(row[39 + k]) for k, n in enumerate(ERANDA)},
        }
    n = {s: len(v) for s, v in out.items()}
    if n != {"男女計": 17, "男": 17, "女": 17}:
        sys.exit(f"【止まりました】第12表の産業の数が想定（16＋計）と違います：{n}")
    return out


# ---- 2. 賃金構造 職種 第5表 ----
def parse_t5(path, names):
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    if "男女計" not in clean(rows[5][3]) or "企業規模計" not in clean(rows[7][3]):
        sys.exit("【止まりました】賃金構造 第5表の形が想定と違います（男女計・企業規模計）")
    out, cur = {}, None
    for r in rows[11:]:
        a = r[1] if len(r) > 1 else None
        if not isinstance(a, str) or not a.strip():
            continue
        vals = dict(zip(FIELDS, [num(v) for v in r[3:11]]))
        if a.startswith("　"):  # 年齢階級の行
            if cur is None:
                continue
            out[cur]["年齢階級"][age_label(a)] = {**vals, **derive(vals)}
        else:
            cur = a.strip()
            out[cur] = {"年齢計": {**vals, **derive(vals)}, "年齢階級": {}}
    miss = [n for n in names if n not in out]
    if miss:
        sys.exit(f"【止まりました】第5表に職種の名前がありません：{miss[:5]}")
    return out


# ---- 3. 賃金構造 産業大分類 第1表 ----
def parse_sangyo(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    out = {}
    for ws in wb.worksheets:
        name = ws.title
        if name.startswith("(民＋公)"):
            continue  # 民営だけ（職種の表と同じ範囲）
        rows = list(ws.iter_rows(min_row=1, max_row=40, values_only=True))
        if clean(rows[5][3]) != "民営事業所" or "企業規模計" not in clean(rows[8][3]):
            sys.exit(f"【止まりました】賃金構造 産業大分類 第1表の {name} の形が想定と違います")
        # 産業は記号（Ｄなど）で引く（表ごとに産業名のかっこの全角・半角が違うため）。産業計は「産業計」
        key = name if name == "産業計" else name[0]
        rec = None
        for r in rows[12:]:
            lab = clean(r[2]) if len(r) > 2 and r[2] is not None else ""
            if "男女計" in lab and "学歴計" in lab:
                vals = dict(zip(FIELDS, [num(v) for v in r[3:11]]))
                rec = {"表の名前": name if name == "産業計" else name[1:], "年齢計": {**vals, **derive(vals)}, "年齢階級": {}}
                continue
            if rec is None:
                continue
            if not re.search(r"歳", lab):
                break  # 学歴（中学など）の段に入ったら終わり
            vals = dict(zip(FIELDS, [num(v) for v in r[3:11]]))
            rec["年齢階級"][age_label(lab)] = {**vals, **derive(vals)}
        if rec is None:
            sys.exit(f"【止まりました】産業大分類 第1表の {name} に男女計・学歴計の行がありません")
        out[key] = rec
    if "産業計" not in out or len(out) != 17:
        sys.exit(f"【止まりました】産業大分類 第1表の産業の数が想定（16＋計）と違います：{len(out)}")
    return out


def run(force=False, register=False):
    if os.path.exists(OUT) and not force:
        got = json.load(open(OUT, encoding="utf-8"))["meta"].get("取得日", "")
        if got[:7] == TODAY[:7]:
            print(f"同じ月（{got}）に取った物があります。取らずに使います：{os.path.relpath(OUT, ROOT)}")
            return
    sh = json.load(open(os.path.join(ROOT, "data", "job", "shokushu_r7.json"), encoding="utf-8"))["data"]["男女計"]
    names = [n for n in sh if n != "不詳"]
    src = {}

    i12 = find_file("00450073", "前の勤め先を辞めた理由", ["年計", "結果表", "2025年"], "12_産業(大分類)")
    p12 = download(i12, "koyodoko_r7_t12_yameta_riyu.xls")
    src["koyodoko_r7_t12"] = i12
    t12 = parse_t12(p12)

    i5 = find_file("00450091", "職種", ["令和７年", "一般労働者", "職種"], "5_職種（小分類）、年齢階級別")
    p5 = download(i5, "chinkin_r7_shokushu_t5_nenrei.xlsx")
    src["chinkin_r7_shokushu_t5"] = i5
    t5 = parse_t5(p5, names)

    i1 = find_file("00450091", "産業大分類", ["令和７年", "一般労働者", "産業大分類"], "1_学歴、年齢階級別きまって支給する現金給与額")
    p1 = download(i1, "chinkin_r7_sangyo_t1.xlsx")
    src["chinkin_r7_sangyo_t1"] = i1
    sg = parse_sangyo(p1)

    # 照らし合わせ：第5表の年齢計が、職種 第1表（shokushu_r7.json）と同じ値か
    bad = [n for n in names if any(t5[n]["年齢計"][k] != sh[n][k] for k in FIELDS)]
    if bad:
        sys.exit(f"【止まりました】第5表の年齢計が職種 第1表と違います：{bad[:5]}")
    # 照らし合わせ：産業計の労働者数が、職種の労働者数の合計と同じか（同じ範囲の表か）
    tot = sum(sh[n]["労働者数"] or 0 for n in sh)
    if abs(tot - sg["産業計"]["年齢計"]["労働者数"]) > 0.5:
        sys.exit(f"【止まりました】産業計の労働者数（{sg['産業計']['年齢計']['労働者数']}）と職種の合計（{tot}）が違います")

    # 「〇〇の40年」の数え（手順書 §3）：20〜24歳から60〜64歳まで5歳ごとの年収が欠けずに取れる職業
    full40 = [n for n in names if all(t5[n]["年齢階級"].get(a, {}).get("年収_千円") is not None for a in AGE9)]

    meta = {"取得日": TODAY, "e_stat_note": estat.NOTE,
            "決まり": "値はダウンロードした表の値そのまま。年収・時給は fetch_job_data.derive と同じ式。順位・差・言い直しは a1_facts.py・a3_facts.py・c_facts.py で計算する",
            "単位": {**dict(zip(FIELDS, UNITS)), "雇用動向調査": "人数は千人、構成比は%"},
            "雇用動向調査の注意": "第12表は「就業経験有りで調査時在籍者で前職雇用者」の入職者（転職入職者）。産業は入職した先（今の勤め先）の産業。就業形態計・年齢計の段",
            "賃金構造の範囲": "一般労働者・民営事業所・企業規模10人以上・男女計（職種の表と同じ範囲。産業計の労働者数が職種の合計と一致することを確かめた）",
            "〇〇の40年": {"条件": "20〜24歳から60〜64歳まで5歳ごとの年収が全部ある職業（不詳を除く）", "数": len(full40),
                         "型にできるか": "10職業以上なら型にできる（手順書 §3）" if len(full40) >= 10 else "10職業に届かない"}}
    data = {"離職理由_産業別": t12, "職種_年齢階級": t5, "産業別_平均": sg, "〇〇の40年の職業": full40}
    json.dump({"meta": meta, "sources": src, "data": data}, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    out_sha = sha256(OUT)
    print(f"保存しました：{os.path.relpath(OUT, ROOT)}（産業 {len(t12['男女計'])}・職種 {len(t5)}・産業の平均 {len(sg)}・〇〇の40年 {len(full40)}職業）")

    sums = {v["file"]: v["sha256"] for v in src.values()}
    sums[os.path.relpath(OUT, ROOT).replace("\\", "/")] = out_sha
    if register:
        # 出典と照合用のハッシュ（手順書 §8・設計書 §4）。書く直前に読み直し、自分のキーだけ足す（ほかのキーは消さない）
        sp = os.path.join(ROOT, "data", "sources.json")
        sources = json.load(open(sp, encoding="utf-8"))
        sources.update(src)
        json.dump(sources, open(sp, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        cp = os.path.join(ROOT, "data", "checksums.json")
        cs = json.load(open(cp, encoding="utf-8"))
        cs.update(sums)
        json.dump(cs, open(cp, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print("sources.json と checksums.json に足しました")
    else:
        print("--register なし。sources.json に足すキー：")
        print(json.dumps(src, ensure_ascii=False, indent=1))
        print("checksums.json に足すキー：")
        print(json.dumps(sums, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--force", action="store_true")
    p.add_argument("--register", action="store_true")
    a = p.parse_args()
    run(a.force, a.register)
