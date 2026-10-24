# 型D「日本人の給料はなぜ30年上がらないのか」の統計を取る（手順書 §3 型D・§8、編集方針の出典の決まり）
#   python -X utf8 pipeline\fetch_d_wage30.py   … 取って、原本を data\raw\ に、読んだ値を data\job\d_wage30.json に保存
# ・取る物（5人の容疑者を照らす統計。どれも「年ごとの値」をできるだけ長く取る）
#   法人企業統計調査（財務省）時系列 金融業・保険業以外 全産業・全規模 … e-Stat API
#   労働力調査（総務省）長期時系列 表9-1 年齢階級・雇用形態別雇用者数 … e-Stat のファイル
#   毎月勤労統計調査（厚生労働省）長期時系列表 賃金指数（就業形態計・一般労働者）と実質賃金指数 … e-Stat のファイル
#   消費者物価指数 2020年基準（総務省）全国 年平均 … e-Stat API
#   家計調査（総務省）家計収支編 二人以上の世帯のうち勤労者世帯 年平均 … e-Stat API
#   賃金構造基本統計調査（厚生労働省）一般労働者 年齢階級別 … e-Stat API（2020年に調べ方が変わったので2019年まで）
#   国民経済計算 2024年度年次推計（内閣府）… 内閣府のサイトの Excel（e-Stat には最新の表がないため。e-Stat の統計に数えない）
# ・ここでは表の値を年ごとに並べるだけ。割合・伸び率・点数は d_facts.py で計算する（Claude は計算しない）
# ・ファイルは data\raw\estat\ と data\raw\esri\ に保存し、URL・取得日を data\sources.json に、ハッシュを data\checksums.json に足す
#   （ほかのキーは消さない）
import datetime, hashlib, json, os, re, sys
import openpyxl, requests, xlrd

sys.path.insert(0, os.path.dirname(__file__))
import estat

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RAW_E = os.path.join(ROOT, "data", "raw", "estat")
RAW_C = os.path.join(ROOT, "data", "raw", "esri")
OUT = os.path.join(ROOT, "data", "job", "d_wage30.json")
TODAY = datetime.date.today().isoformat()
DBVIEW = "https://www.e-stat.go.jp/dbview?sid={}"
FROM = 1994  # 1995年と比べるため、その1年前から取る

# 内閣府 国民経済計算 2024年度年次推計（2020年基準・2008SNA）
ESRI_TOP = "https://www.esri.cao.go.jp/jp/sna/data/data_list/kakuhou/files/2024/2024_kaku_top.html"
ESRI_TBL = "https://www.esri.cao.go.jp/jp/sna/data/data_list/kakuhou/files/2024/tables/{}"


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        h.update(f.read())
    return h.hexdigest()


def rel(path):
    return os.path.relpath(path, ROOT).replace("\\", "/")


def save_json(obj, path):
    text = json.dumps(obj, ensure_ascii=False, indent=1)
    if estat._app_id() in text:  # キーが原本に紛れていないか（値は出さない）
        sys.exit("【止まりました】保存する原本にキーが入っています。保存をやめました")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)


def get(url, path):
    r = requests.get(url, timeout=180, headers={"User-Agent": "Mozilla/5.0"})
    if r.status_code != 200 or len(r.content) < 1000:
        sys.exit(f"【止まりました】{os.path.basename(path)} を取れませんでした（{r.status_code}）")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as f:
        f.write(r.content)
    return r


# ---- e-Stat API の表 ----
def api_table(sid, name, **params):
    """getStatsData を取り、原本の JSON を保存して、値の一覧と表の情報を返す"""
    d = estat.call("getStatsData", statsDataId=sid, limit=100000, **params)
    sd = d["GET_STATS_DATA"]
    if sd["RESULT"]["STATUS"] != 0:
        sys.exit(f"【止まりました】{sid}：{sd['RESULT']['ERROR_MSG']}")
    path = os.path.join(RAW_E, name)
    save_json(d, path)
    ti = sd["STATISTICAL_DATA"]["TABLE_INF"]
    title = ti["TITLE"]["$"] if isinstance(ti["TITLE"], dict) else ti["TITLE"]
    vals = sd["STATISTICAL_DATA"]["DATA_INF"]["VALUE"]
    vals = [vals] if isinstance(vals, dict) else vals
    info = {"statsDataId": sid, "統計名": ti["STATISTICS_NAME"], "府省": ti["GOV_ORG"]["$"], "表名": title,
            "公表日": ti.get("OPEN_DATE"), "更新日": ti.get("UPDATED_DATE"), "url": DBVIEW.format(sid),
            "取り方": "e-Stat API（getStatsData）", "条件": params, "e_stat": True,
            "file": rel(path), "sha256": sha256(path), "fetched": TODAY}
    return vals, info


def fnum(s):
    try:
        return float(s)
    except (TypeError, ValueError):
        return None  # 「-」「…」など、値のない欄


# ---- e-Stat のファイル ----
def catalog(stats_code, word, stat_inf_id):
    """カタログ API で statInfId の資源を探し、表名・公表日・URL を返す"""
    start = 1
    while True:
        d = estat.call("getDataCatalog", statsCode=stats_code, searchWord=word, limit=100, startPosition=start)
        r = d["GET_DATA_CATALOG"]
        ds = r.get("DATA_CATALOG_LIST_INF", {}).get("DATA_CATALOG_INF", [])
        ds = [ds] if isinstance(ds, dict) else ds
        for x in ds:
            res = x["RESOURCES"]["RESOURCE"]
            res = [res] if isinstance(res, dict) else res
            for rr in res:
                if f"statInfId={stat_inf_id}" in rr["URL"]:
                    return {"statInfId": stat_inf_id, "dataset": x["DATASET"]["TITLE"].get("NAME"),
                            "表名": rr["TITLE"].get("NAME"), "公表日": rr.get("RELEASE_DATE"),
                            "更新日": rr.get("LAST_MODIFIED_DATE"), "url": x["DATASET"].get("LANDING_PAGE"),
                            "download": rr["URL"], "licence": rr.get("RESOURCE_LICENCE_ID"),
                            "取り方": "e-Stat のファイル", "e_stat": True}
        nxt = r.get("DATA_CATALOG_LIST_INF", {}).get("RESULT_INF", {}).get("NEXT_KEY")
        if not nxt:
            sys.exit(f"【止まりました】e-Stat のカタログに {stat_inf_id} が見つかりません")
        start = nxt


def estat_file(stats_code, word, stat_inf_id, name, gov, stat_name):
    info = {"統計名": stat_name, "府省": gov, **catalog(stats_code, word, stat_inf_id)}
    path = os.path.join(RAW_E, name)
    get(info["download"], path)
    info.update(file=rel(path), sha256=sha256(path), fetched=TODAY)
    return path, info


# ---- 1. 法人企業統計調査 ----
HOJIN_SID = "0003060791"
HOJIN_ITEMS = {"002": "現金・預金", "226": "利益剰余金", "066": "従業員給与", "235": "従業員賞与",
               "065": "役員給与", "057": "役員賞与", "067": "福利厚生費", "073": "付加価値", "072": "期中平均従業員数"}


def hojin():
    vals, info = api_table(HOJIN_SID, "d_wage30_hojin.json", cdCat01=",".join(HOJIN_ITEMS), cdCat02="104",
                           cdCat03="26", cdTimeFrom=f"{FROM}0")
    info["範囲"] = "金融業、保険業以外の業種・全産業・全規模・年度"
    out = {v: {} for v in HOJIN_ITEMS.values()}
    for x in vals:
        y = int(x["@time"][:4])
        out[HOJIN_ITEMS[x["@cat01"]]][y] = fnum(x["$"])
    units = {v: ("人" if k == "072" else "百万円") for k, v in HOJIN_ITEMS.items()}
    return out, units, info


# ---- 2. 労働力調査 長期時系列 表9-1 ----
LFS_ID = "000021915916"
AGES = {"15～24歳": 17, "25～34歳": 45, "35～44歳": 59, "45～54歳": 73, "55～64歳": 87, "65歳以上": 101}


def lfs():
    path, info = estat_file("00200531", "長期時系列", LFS_ID, "d_wage30_lfs_9-1.xlsx", "総務省", "労働力調査")
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True)["男女計"]
    rows = list(ws.iter_rows(values_only=True))
    # 見出しの列の名前を確かめる（列の位置がずれていたら止める。セルの中の改行は除いて比べる）
    def head(c):
        return "".join(re.sub(r"\s", "", str(rows[i][c])) for i in range(13) if c < len(rows[i]) and rows[i][c] is not None)
    if "役員を除く雇用者" not in head(5) or "非正規の職員・従業員" not in head(7) or "非正規の職員・従業員" not in head(16):
        sys.exit("【止まりました】労働力調査の表の総数の列の位置が想定と違います")
    for name, c in AGES.items():
        if name not in head(c) or "就業者" not in head(c) or "役員を除く雇用者" not in head(c + 2) or "非正規の職員・従業員" not in head(c + 4):
            sys.exit(f"【止まりました】労働力調査の {name} の列の位置が想定と違います")
    out = {}
    for r in rows:
        lab = str(r[2]) if len(r) > 2 and r[2] is not None else ""
        m1 = re.fullmatch(r"Feb\. (\d{4})", lab.strip())
        m2 = re.fullmatch(r"(\d{4}) Jan\.-Mar\. average", lab.strip())
        if not (m1 or m2):
            continue
        key = f"{m1.group(1)}年2月" if m1 else f"{m2.group(1)}年1～3月"
        rec = {"役員を除く雇用者": fnum(r[5]), "正規": fnum(r[6]), "非正規": fnum(r[7]), "非正規の割合": fnum(r[16])}
        for name, c in AGES.items():
            rec[f"{name}_役員を除く雇用者"] = fnum(r[c + 2])
            rec[f"{name}_非正規"] = fnum(r[c + 4])
        out[key] = rec
    info["範囲"] = "男女計。2001年までは2月（労働力調査特別調査）、2002年からは四半期平均（詳細集計）。ここでは各年の2月と1～3月平均を使う"
    info["注意"] = "特別調査（2月）と詳細集計（四半期）は調べ方が違う（表の注記）。季節をそろえるため、2001年までは2月、2002年からは1～3月平均を使う"
    return out, info


# ---- 3. 毎月勤労統計調査 長期時系列表 ----
MAIKIN = {"名目賃金指数_就業形態計": ("000032189720", "d_wage30_maikin_7_kei.xls", "就業形態計"),
          "名目賃金指数_一般労働者": ("000032189722", "d_wage30_maikin_9_ippan.xls", "一般労働者"),
          "実質賃金指数_就業形態計": ("000032189738", "d_wage30_maikin_25-1_real.xls", "就業形態計")}


def maikin():
    out, infos = {}, {}
    for key, (sid, name, form) in MAIKIN.items():
        path, info = estat_file("00450071", "時系列表", sid, name, "厚生労働省", "毎月勤労統計調査")
        ws = xlrd.open_workbook(path).sheet_by_name("TL")
        top = " ".join(str(ws.row_values(i)[5]) for i in range(4))
        if form not in top or "調査産業計" not in top or "５人以上" not in top or "年平均" not in " ".join(map(str, ws.row_values(7))):
            sys.exit(f"【止まりました】毎月勤労統計の {name} の表の形が想定と違います")
        if "令和２年平均＝１００" not in " ".join(map(str, ws.row_values(5))):
            sys.exit(f"【止まりました】毎月勤労統計の {name} の基準年が2020年ではありません")
        ser = {}
        for i in range(9, ws.nrows):
            r = ws.row_values(i)
            if r[0] == "" or str(r[0]).startswith("　"):
                break  # 指数の段の終わり（下は増減率の段）
            y, v = int(r[0]), fnum(r[1])
            if y >= FROM and v is not None:
                ser[y] = v
        out[key] = ser
        info["範囲"] = "事業所規模5人以上・調査産業計・年平均・2020年平均＝100"
        infos[key] = info
    infos["実質賃金指数_就業形態計"]["注意"] = "消費者物価指数（持家の帰属家賃を除く総合）で実質化した指数（表の題）"
    return out, infos


# ---- 4. 消費者物価指数 ----
def cpi():
    vals, info = api_table("0003427113", "d_wage30_cpi.json", cdTab="1", cdCat01="0001,0163", cdArea="00000")
    names = {"0001": "総合", "0163": "持家の帰属家賃を除く総合"}
    out = {v: {} for v in names.values()}
    for x in vals:
        t = x["@time"]
        if t.endswith("000000") and len(t) == 10 and int(t[:4]) >= FROM:
            out[names[x["@cat01"]]][int(t[:4])] = fnum(x["$"])
    info["範囲"] = "全国・年平均・2020年＝100"
    return out, info


# ---- 5. 家計調査 勤労者世帯 ----
KAKEI_ITEMS = {"019": "実収入", "193": "非消費支出", "194": "直接税", "198": "社会保険料"}
KAKEI_SET = {"02": "勤労者世帯（農林漁家世帯を除く）", "04": "勤労者世帯（2000年～）"}


def kakei():
    vals, info = api_table("0002070003", "d_wage30_kakei.json", cdCat01=",".join(KAKEI_ITEMS),
                           cdCat02=",".join(KAKEI_SET), cdArea="00000")
    out = {s: {v: {} for v in KAKEI_ITEMS.values()} for s in KAKEI_SET.values()}
    for x in vals:
        v = fnum(x["$"])
        if v is not None and int(x["@time"][:4]) >= FROM:
            out[KAKEI_SET[x["@cat02"]]][KAKEI_ITEMS[x["@cat01"]]][int(x["@time"][:4])] = v
    info["範囲"] = "二人以上の世帯のうち勤労者世帯・全国・年平均・1世帯当たり1か月（円）"
    info["注意"] = "世帯の区分が2000年から『勤労者世帯』、1999年までは『勤労者世帯（農林漁家世帯を除く）』。重なる年の割合を d_facts.py で照らしてからつなぐ"
    return out, info


# ---- 6. 賃金構造基本統計調査 一般労働者 年齢階級別 ----
AGE5 = ["20～24歳", "25～29歳", "30～34歳", "35～39歳", "40～44歳", "45～49歳", "50～54歳", "55～59歳", "60～64歳"]


def chinkin():
    # 1999～2003年（産業計（民・公営計）・企業規模計・学歴計・男女計・労働者）
    v1, i1 = api_table("0003130256", "d_wage30_chinkin_1999-2003.json", cdTab="42,45", cdCat02="01",
                       cdCat05="01", cdCat06="10", cdCat07="01", cdCat08="10")
    a1 = {"nenrei01": "年齢計", "nenrei04": "20～24歳", "nenrei05": "25～29歳", "nenrei06": "30～34歳",
          "nenrei07": "35～39歳", "nenrei08": "40～44歳", "nenrei09": "45～49歳", "nenrei10": "50～54歳",
          "nenrei11": "55～59歳", "nenrei12": "60～64歳"}
    # 2010～2019年（産業計・企業規模計（10人以上）・学歴計・男女計・民営＋公営）
    v2, i2 = api_table("0003084009", "d_wage30_chinkin_2010-2019.json", cdTab="42,45", cdCat01="01",
                       cdCat02="01", cdCat03="01", cdCat04="01", cdCat06="01")
    a2 = {"01": "年齢計", "03": "20～24歳", "04": "25～29歳", "05": "30～34歳", "06": "35～39歳", "07": "40～44歳",
          "08": "45～49歳", "09": "50～54歳", "10": "55～59歳", "11": "60～64歳"}
    tab = {"42": "所定内給与額", "45": "労働者数"}
    out = {"所定内給与額": {}, "労働者数": {}}
    for vals, ages, akey in ((v1, a1, "@cat04"), (v2, a2, "@cat05")):
        for x in vals:
            if x[akey] not in ages:
                continue
            v = fnum(x["$"])
            if v is not None:
                out[tab[x["@tab"]]].setdefault(int(x["@time"][:4]), {})[ages[x[akey]]] = v
    for k in out:  # 年の並びを数字の順に
        out[k] = dict(sorted(out[k].items()))
    i1["範囲"] = "一般労働者・産業計（民・公営計）・企業規模計・学歴計・男女計"
    i2["範囲"] = "一般労働者・産業計・企業規模計（10人以上）・学歴計・男女計・民営＋公営"
    note = "2020年から調べ方が変わったので、年齢をそろえた比べは2019年までで行う。単位は所定内給与額が千円、労働者数が十人"
    i1["注意"] = i2["注意"] = note
    return out, i1, i2


# ---- 7. 国民経済計算（内閣府） ----
def esri_file(fname, key, title):
    path = os.path.join(RAW_C, f"d_wage30_{fname}")
    r = get(ESRI_TBL.format(fname), path)
    info = {"統計名": "国民経済計算（GDP統計）", "府省": "内閣府", "表名": title, "url": ESRI_TOP,
            "download": ESRI_TBL.format(fname), "サーバーの更新日時": r.headers.get("Last-Modified"),
            "公表日": None, "取り方": "内閣府のサイトの Excel", "e_stat": False,
            "e_stat_の判断": "e-Stat には2024年度年次推計の表がない（e-Stat の国民経済計算は古い基準の表だけ）。e-Stat の統計に数えない",
            "file": rel(path), "sha256": sha256(path), "fetched": TODAY}
    return path, info


def sna_rows(path, sheet, want):
    """暦年の表から、項目名に want の文字を含む行の値を {年: 値} で返す"""
    ws = openpyxl.load_workbook(path, data_only=True, read_only=True)[sheet]
    rows = list(ws.iter_rows(values_only=True))
    years = [int(str(c)) if c is not None and str(c).isdigit() else None for c in rows[6]]
    out = {}
    for name, pat in want.items():
        hit = [r for r in rows if r and isinstance(r[0], str) and re.sub(r"\s", "", r[0]).startswith(pat)]
        if len(hit) != 1:
            sys.exit(f"【止まりました】国民経済計算の {sheet} に「{pat}」の行が {len(hit)} 本あります")
        out[name] = {y: fnum(v) for y, v in zip(years, hit[0]) if y is not None and fnum(v) is not None}
    return out


def sna():
    out, infos = {}, {}
    p, i = esri_file("2024fcm1rn_jp.xlsx", "real", "2024年度国民経済計算 主要系列表 1.国内総生産（支出側） 実質（2020暦年連鎖価格） 暦年")
    out.update({"実質GDP": sna_rows(p, "実数", {"x": "５．国内総生産（支出側）"})["x"]})
    infos["sna_gdp_real"] = i
    p, i = esri_file("2024fcm1n_jp.xlsx", "nominal", "2024年度国民経済計算 主要系列表 1.国内総生産（支出側） 名目 暦年")
    out.update({"名目GDP": sna_rows(p, "実数", {"x": "５．国内総生産（支出側）"})["x"]})
    infos["sna_gdp_nominal"] = i
    p, i = esri_file("2024fcm2_jp.xlsx", "income", "2024年度国民経済計算 主要系列表 2.国民所得・国民可処分所得の分配 暦年")
    out.update(sna_rows(p, "実数", {"雇用者報酬": "１．雇用者報酬", "賃金・俸給": "（１）賃金・俸給", "雇主の社会負担": "（２）雇主の社会負担"}))
    infos["sna_income"] = i
    p, i = esri_file("2024s3_jp.xlsx", "labor", "2024年度国民経済計算 付表3 経済活動別の就業者数・雇用者数、労働時間数 暦年")
    for sheet, name in (("暦年（１）就業者", "就業者数"), ("暦年（２）雇用者", "雇用者数"), ("暦年（３）労働時間数", "雇用者1人当たり労働時間")):
        ws = openpyxl.load_workbook(p, data_only=True, read_only=True)[sheet]
        rows = list(ws.iter_rows(values_only=True))
        years = [c if isinstance(c, int) else None for c in rows[6]]
        tot = [r for r in rows if r and isinstance(r[0], str) and re.sub(r"\s", "", r[0]) == "合計"]
        if not tot:
            sys.exit(f"【止まりました】国民経済計算の {sheet} に合計の行がありません")
        out[name] = {y: fnum(v) for y, v in zip(years, tot[0]) if y is not None and fnum(v) is not None}
    infos["sna_labor"] = i
    units = {"実質GDP": "10億円（2020暦年連鎖価格）", "名目GDP": "10億円", "雇用者報酬": "10億円", "賃金・俸給": "10億円",
             "雇主の社会負担": "10億円", "就業者数": "万人", "雇用者数": "万人", "雇用者1人当たり労働時間": "時間（年）"}
    for k in infos.values():
        k["範囲"] = "暦年・2020年基準（2008SNA）"
    return out, units, infos


def span(d):
    ys = sorted(int(y) for y in d)
    return f"{ys[0]}～{ys[-1]}" if ys else "なし"


def run():
    src, series = {}, {}

    h, hu, hi = hojin()
    src["d_wage30_hojin"] = hi
    for k, v in h.items():
        series[f"法人企業統計_{k}"] = {"source": "d_wage30_hojin", "単位": hu[k], "年の区切り": "年度", "年の範囲": span(v), "値": v}
    # 2006年度までは従業員給与に賞与が入り、2007年度から従業員賞与が別の項目になった（推測。値の段差で確かめる）
    lf, li = lfs()
    src["d_wage30_lfs"] = li
    series["労働力調査_表9-1"] = {"source": "d_wage30_lfs", "単位": "万人（割合は%）", "年の区切り": "2月または1～3月平均",
                                "年の範囲": f"{min(lf)[:4]}～{max(lf)[:4]}", "値": lf}
    mk, mi = maikin()
    for k, v in mk.items():
        skey = f"d_wage30_maikin_{k}"
        src[skey] = mi[k]
        series[f"毎月勤労統計_{k}"] = {"source": skey, "単位": "指数（2020年＝100）", "年の区切り": "暦年", "年の範囲": span(v), "値": v}
    cp, ci = cpi()
    src["d_wage30_cpi"] = ci
    for k, v in cp.items():
        series[f"消費者物価指数_{k}"] = {"source": "d_wage30_cpi", "単位": "指数（2020年＝100）", "年の区切り": "暦年", "年の範囲": span(v), "値": v}
    kk, ki = kakei()
    src["d_wage30_kakei"] = ki
    for s, items in kk.items():
        for k, v in items.items():
            series[f"家計調査_{s}_{k}"] = {"source": "d_wage30_kakei", "単位": "円（1世帯当たり1か月）", "年の区切り": "暦年", "年の範囲": span(v), "値": v}
    ch, c1, c2 = chinkin()
    src["d_wage30_chinkin_1999"] = c1
    src["d_wage30_chinkin_2010"] = c2
    for k, v in ch.items():
        series[f"賃金構造_{k}"] = {"source": "d_wage30_chinkin_1999 / d_wage30_chinkin_2010", "単位": "千円" if k == "所定内給与額" else "十人",
                                  "年の区切り": "暦年（6月分）", "年の範囲": span(v), "値": v}
    sn, su, si = sna()
    src.update({f"d_wage30_{k}": v for k, v in si.items()})
    skey = {"実質GDP": "d_wage30_sna_gdp_real", "名目GDP": "d_wage30_sna_gdp_nominal", "雇用者報酬": "d_wage30_sna_income",
            "賃金・俸給": "d_wage30_sna_income", "雇主の社会負担": "d_wage30_sna_income", "就業者数": "d_wage30_sna_labor",
            "雇用者数": "d_wage30_sna_labor", "雇用者1人当たり労働時間": "d_wage30_sna_labor"}
    for k, v in sn.items():
        series[f"国民経済計算_{k}"] = {"source": skey[k], "単位": su[k], "年の区切り": "暦年", "年の範囲": span(v), "値": v}

    # 表の名前・URL・取得日を各系列にも付ける（facts.json の出典カードに使う）
    for s in series.values():
        k = s["source"].split(" / ")[0]
        s.update({"表名": src[k]["表名"], "URL": src[k]["url"], "取得日": src[k]["fetched"], "e_stat": src[k]["e_stat"]})

    meta = {"題": "日本人の給料はなぜ30年上がらないのか（型D）", "取得日": TODAY,
            "e_stat_note": estat.NOTE, "決まり": "値はダウンロードした表の値そのまま。割合・伸び率・点数は d_facts.py で計算する",
            "e-Stat API で取った表": [k for k, v in src.items() if v["取り方"].startswith("e-Stat API")],
            "e-Stat にない表": [k for k, v in src.items() if not v["e_stat"]]}
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    json.dump({"meta": meta, "series": series}, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    # 出典と照合用のハッシュ（手順書 §8・設計書 §4）。ほかのキーは消さない
    sp = os.path.join(ROOT, "data", "sources.json")
    sources = json.load(open(sp, encoding="utf-8")) if os.path.exists(sp) else {}
    sources.update(src)
    json.dump(sources, open(sp, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    cpth = os.path.join(ROOT, "data", "checksums.json")
    sums = json.load(open(cpth, encoding="utf-8")) if os.path.exists(cpth) else {}
    for k in src.values():
        sums[k["file"]] = k["sha256"]
    sums[rel(OUT)] = sha256(OUT)
    json.dump(sums, open(cpth, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"系列 {len(series)}・表 {len(src)} を保存しました：{rel(OUT)}")
    for k, s in series.items():
        print(f"  {k}：{s['年の範囲']}")


if __name__ == "__main__":
    run()
